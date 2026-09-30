"""Excel での一括登録・出力。

テンプレート (部材 / 仕入先 シート) をダウンロード → 記入 → アップロード → プレビューで確認 → 反映。
部材は品番、仕入先は社名で照合し、既にあれば更新、無ければ新規登録する。
「部材一覧を Excel で出力」したファイルをそのまま編集して取り込み直すこともできる。
"""
import io
import json
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import db

# (見出し, 項目キー, 必須, 説明)
MATERIAL_COLS = [
    ("品番", "part_number", True, "部材を識別するキー。既にある品番は上書き更新"),
    ("品名", "name", True, ""),
    ("メーカー", "maker", False, ""),
    ("仕様", "spec", False, ""),
    ("数量", "quantity", False, "必要数量 (整数・省略時 1)"),
    ("単位", "unit", False, "省略時「個」"),
    ("必要納期", "required_date", False, "日付 (2026/12/31 など)"),
    ("予算単価", "budget_unit_price", False, "円・税抜"),
    ("特注品", "custom_item", False, "はい / いいえ"),
    ("担当者", "owner", False, "チーム・仕入先 画面に登録済みの氏名かメール"),
    ("主仕入先", "supplier", False, "仕入先シートか登録済みの社名"),
    ("商品ページURL", "watch_urls", False, "ネットで価格を取得するページ。複数は改行で区切る"),
    ("確認周期(日)", "confirm_interval_days", False, "メールで価格・納期を確認する周期。空欄=既定、0=しない"),
    ("自動確認メール", "auto_confirm", False, "はい = 確認期限に確認メールを自動送信"),
    ("備考", "notes", False, ""),
    ("現在の単価", "current_price", False, "わかっていれば記入 (価格の履歴に記録)"),
    ("現在の納期(日)", "current_lead", False, "わかっていれば記入"),
]
SUPPLIER_COLS = [
    ("社名", "name", True, "既にある社名は上書き更新"),
    ("ご担当者名", "contact_name", False, ""),
    ("メール", "email", True, "見積依頼・確認・注文の送り先"),
]
EXPORT_EXTRA = ["最新単価(参考)", "最新納期(参考)", "取得日(参考)", "取得元(参考)"]
FIELD_LABELS = {k: h for h, k, _, _ in MATERIAL_COLS}

HEAD_FILL = PatternFill("solid", fgColor="2F5BEA")
REQ_FILL = PatternFill("solid", fgColor="B42318")
REF_FILL = PatternFill("solid", fgColor="667085")


def _sheet(ws, cols, extra=()):
    for i, (h, _, req, note) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = REQ_FILL if req else HEAD_FILL
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(36, len(h) * 2 + 6))
    for j, h in enumerate(extra, len(cols) + 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = REF_FILL
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.freeze_panes = "B2"


def _yes_no_validation(ws, cols, keys, rows=2000):
    dv = DataValidation(type="list", formula1='"はい,いいえ"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, (_, k, _, _) in enumerate(cols, 1):
        if k in keys:
            dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}{rows}")


def template():
    wb = Workbook()
    ws = wb.active
    ws.title = "部材"
    _sheet(ws, MATERIAL_COLS)
    ws.append(["SFJ10-100", "リニアシャフト φ10", "ミスミ", "L=100", 20, "本", "2026/12/15", 900, "いいえ", "",
               "株式会社部品商事", "https://jp.misumi-ec.com/vona2/detail/...", "", "いいえ", "記入例 (この行は消してください)",
               "", ""])
    _yes_no_validation(ws, MATERIAL_COLS, {"custom_item", "auto_confirm"})
    ws2 = wb.create_sheet("仕入先")
    _sheet(ws2, SUPPLIER_COLS)
    ws2.append(["株式会社部品商事", "鈴木", "sales@buhin.example.com"])
    ws3 = wb.create_sheet("説明")
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 90
    ws3.append(["部材トラッカー 一括登録テンプレート"])
    ws3["A1"].font = Font(bold=True, size=14)
    ws3.append([])
    ws3.append(["使い方", "「部材」「仕入先」シートに記入して保存し、アプリの「Excel 一括登録」画面からアップロードします。"])
    ws3.append(["", "取り込む前に、新規・更新・エラーの一覧を確認できます。赤い見出しは必須項目です。"])
    ws3.append(["", "品番が既にある部材は、記入した項目だけ上書きします (空欄の項目は変更しません)。"])
    ws3.append([])
    ws3.append(["項目", "説明"])
    for h, _, req, note in MATERIAL_COLS:
        ws3.append([h + (" *" if req else ""), note])
    ws3.append([])
    for h, _, req, note in SUPPLIER_COLS:
        ws3.append([f"仕入先: {h}" + (" *" if req else ""), note])
    return _save(wb)


def export_materials(conn):
    wb = Workbook()
    ws = wb.active
    ws.title = "部材"
    _sheet(ws, MATERIAL_COLS, EXPORT_EXTRA)
    staff = {r["id"]: r["name"] for r in conn.execute("SELECT id, name FROM staff")}
    sups = {r["id"]: r["name"] for r in conn.execute("SELECT id, name FROM suppliers")}
    for m in conn.execute("SELECT * FROM materials WHERE active=1 ORDER BY part_number"):
        last = db.latest_observation(conn, m["id"])
        ws.append([m["part_number"], m["name"], m["maker"], m["spec"], m["quantity"], m["unit"],
                   _xl_date(m["required_date"]), m["budget_unit_price"], "はい" if m["custom_item"] else "いいえ",
                   staff.get(m["owner_id"], ""), sups.get(m["preferred_supplier_id"], ""), m["watch_urls"],
                   m["confirm_interval_days"], "はい" if m["auto_confirm"] else "いいえ", m["notes"], None, None,
                   last["unit_price"] if last else None, last["lead_time_days"] if last else None,
                   last["observed_at"][:10] if last else None, (last["vendor"] or last["source"]) if last else None])
    _yes_no_validation(ws, MATERIAL_COLS, {"custom_item", "auto_confirm"})
    ws2 = wb.create_sheet("仕入先")
    _sheet(ws2, SUPPLIER_COLS)
    for s in conn.execute("SELECT * FROM suppliers ORDER BY name"):
        ws2.append([s["name"], s["contact_name"], s["email"]])
    return _save(wb)


def export_rows(title, headers, rows):
    """汎用: 表を 1 シートの Excel にする (履歴・価格推移の出力用)。"""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(48, len(h) * 2 + 6))
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    return _save(wb)


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xl_date(s):
    try:
        return date.fromisoformat(s) if s else None
    except ValueError:
        return s


# ------------------------------------------------------------------ 読み込み・検証
def _cell(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _as_int(v):
    if v is None:
        return None
    return int(float(str(v).replace(",", "")))


def _as_float(v):
    if v is None:
        return None
    return float(str(v).replace(",", "").replace("¥", "").replace("￥", "").replace("円", ""))


def _as_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).replace("年", "/").replace("月", "/").replace("日", "").replace("-", "/").replace(".", "/")
    parts = [p for p in s.split("/") if p]
    y, mth, d = (int(parts[0]), int(parts[1]), int(parts[2])) if len(parts) == 3 else (None, None, None)
    return date(y, mth, d).isoformat()


def _as_bool(v):
    if v is None:
        return None
    return str(v).strip().lower() in ("はい", "yes", "y", "true", "1", "○", "〇", "on")


def _rows(ws, cols):
    headers = [(_cell(c.value) or "") for c in ws[1]]
    idx = {}
    for h, key, _, _ in cols:
        if h in headers:
            idx[key] = headers.index(h)
    missing = [h for h, key, req, _ in cols if req and key not in idx]
    out = []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        vals = {k: _cell(row[i]) if i < len(row) else None for k, i in idx.items()}
        if not any(v is not None for v in vals.values()):
            continue
        out.append((n, vals))
    return out, missing


def parse(conn, data):
    """アップロードされた Excel を読み、反映内容のプレビューを返す。"""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"Excel ファイルとして読めませんでした ({e})") from e
    result = {"suppliers": [], "materials": [], "errors": [], "warnings": []}
    new_sup_names = set()

    if "仕入先" in wb.sheetnames:
        rows, missing = _rows(wb["仕入先"], SUPPLIER_COLS)
        if missing:
            result["errors"].append(f"仕入先シートに列がありません: {', '.join(missing)}")
        for n, v in rows:
            if not v.get("name") or not v.get("email") or "@" not in str(v.get("email")):
                result["suppliers"].append({"row": n, "action": "error", "values": _jsonable(v),
                                            "message": "社名とメール (xx@xx) は必須です"})
                continue
            ex = conn.execute("SELECT * FROM suppliers WHERE name=?", (v["name"],)).fetchone()
            changes = _diff(ex, v, ("contact_name", "email")) if ex else None
            result["suppliers"].append({"row": n, "action": "update" if ex else "new", "id": ex["id"] if ex else None,
                                        "values": _jsonable(v), "changes": changes})
            new_sup_names.add(v["name"])

    ws = wb["部材"] if "部材" in wb.sheetnames else wb.worksheets[0]
    rows, missing = _rows(ws, MATERIAL_COLS)
    if missing:
        result["errors"].append(f"部材シートに必須の列がありません: {', '.join(missing)} (テンプレートの見出しを使ってください)")
        return result
    staff = conn.execute("SELECT * FROM staff WHERE active=1").fetchall()
    seen = set()
    for n, v in rows:
        errs, clean = [], {}
        if not v.get("part_number") or not v.get("name"):
            errs.append("品番と品名は必須です")
        pn = str(v.get("part_number") or "").strip()
        if pn in seen:
            errs.append("同じ品番が Excel 内で重複しています")
        seen.add(pn)
        for key, fn in (("quantity", _as_int), ("budget_unit_price", _as_float), ("confirm_interval_days", _as_int),
                        ("current_price", _as_float), ("current_lead", _as_int), ("required_date", _as_date)):
            try:
                clean[key] = fn(v.get(key))
            except (ValueError, TypeError):
                errs.append(f"「{FIELD_LABELS[key]}」の値を読めません ({v.get(key)})")
        for key in ("custom_item", "auto_confirm"):
            clean[key] = _as_bool(v.get(key))
        for key in ("name", "maker", "spec", "unit", "notes"):
            clean[key] = None if v.get(key) is None else str(v.get(key))
        clean["part_number"] = pn
        if v.get("watch_urls"):
            urls = [u.strip() for u in str(v["watch_urls"]).replace(",", "\n").splitlines() if u.strip()]
            bad = [u for u in urls if not u.startswith(("http://", "https://"))]
            if bad:
                errs.append(f"URL が正しくありません: {bad[0][:40]}")
            clean["watch_urls"] = "\n".join(urls)
        if v.get("owner"):
            o = next((s for s in staff if s["name"] == str(v["owner"]) or (s["email"] or "").lower() == str(v["owner"]).lower()), None)
            if o:
                clean["owner_id"] = o["id"]
            else:
                result["warnings"].append(f"部材 {n} 行目: 担当者「{v['owner']}」が見つからないため空欄にします")
        if v.get("supplier"):
            s = conn.execute("SELECT id FROM suppliers WHERE name=?", (str(v["supplier"]),)).fetchone()
            if s:
                clean["preferred_supplier_id"] = s["id"]
            elif str(v["supplier"]) in new_sup_names:
                clean["supplier_name"] = str(v["supplier"])  # 仕入先シートで新規登録するもの
            else:
                errs.append(f"主仕入先「{v['supplier']}」が未登録です (仕入先シートに追加してください)")
        if errs:
            result["materials"].append({"row": n, "action": "error", "values": _jsonable(v), "message": " / ".join(errs)})
            continue
        ex = conn.execute("SELECT * FROM materials WHERE part_number=?", (pn,)).fetchone()
        fields = {k: val for k, val in clean.items() if val is not None and k not in ("current_price", "current_lead")}
        entry = {"row": n, "action": "update" if ex else "new", "id": ex["id"] if ex else None,
                 "fields": fields, "current_price": clean["current_price"], "current_lead": clean["current_lead"],
                 "values": _jsonable(v)}
        if ex:
            entry["changes"] = _diff(ex, fields, [k for k in fields if k not in ("supplier_name", "part_number")])
            if ex["active"] == 0:
                entry["changes"]["active"] = ["無効", "有効"]
            if not entry["changes"] and clean["current_price"] is None and clean["current_lead"] is None:
                entry["action"] = "same"
        result["materials"].append(entry)
    return result


def _jsonable(v):
    return {k: (x.isoformat() if isinstance(x, (date, datetime)) else x) for k, x in v.items()}


def _diff(row, new, keys):
    out = {}
    for k in keys:
        if k not in new or new[k] is None:
            continue
        old = row[k] if k in row.keys() else None
        nv = new[k]
        if isinstance(nv, bool):
            nv = 1 if nv else 0
        if str(old if old is not None else "") != str(nv):
            out[k] = [old, nv]
    return out


def apply(conn, preview, user_id=None, actor=None, client=None):
    """プレビューの内容を反映する。戻り値: 件数の dict"""
    counts = {"supplier_new": 0, "supplier_update": 0, "material_new": 0, "material_update": 0, "price": 0}
    for s in preview["suppliers"]:
        v = s["values"]
        if s["action"] == "new":
            conn.execute("INSERT INTO suppliers (name, contact_name, email) VALUES (?,?,?)",
                         (v["name"], v.get("contact_name"), v["email"]))
            counts["supplier_new"] += 1
        elif s["action"] == "update" and s.get("changes"):
            conn.execute("UPDATE suppliers SET contact_name=COALESCE(?, contact_name), email=? WHERE id=?",
                         (v.get("contact_name"), v["email"], s["id"]))
            counts["supplier_update"] += 1
    conn.commit()
    if counts["supplier_new"] or counts["supplier_update"]:
        db.log_activity(conn, user_id, "excel_import", None,
                        f"仕入先 新規 {counts['supplier_new']} 件 / 更新 {counts['supplier_update']} 件", actor, client)
    for m in preview["materials"]:
        if m["action"] in ("error", "same") and not (m["action"] == "same" and (m.get("current_price") is not None or m.get("current_lead") is not None)):
            continue
        f = dict(m["fields"])
        if f.pop("supplier_name", None):
            s = conn.execute("SELECT id FROM suppliers WHERE name=?", (m["values"]["supplier"],)).fetchone()
            if s:
                f["preferred_supplier_id"] = s["id"]
        for k in ("custom_item", "auto_confirm"):
            if k in f:
                f[k] = 1 if f[k] else 0
        if m["action"] == "new":
            f.setdefault("quantity", 1)
            f.setdefault("unit", "個")
            cols = list(f) + ["created_by"]
            mid = conn.execute(f"INSERT INTO materials ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                               [f[k] for k in f] + [user_id]).lastrowid
            counts["material_new"] += 1
            detail = "Excel から新規登録"
        else:
            mid = m["id"]
            f.pop("part_number", None)
            if f:
                conn.execute(f"UPDATE materials SET {', '.join(k + '=?' for k in f)}, active=1, updated_by=?, "
                             "updated_at=datetime('now','localtime') WHERE id=?", [f[k] for k in f] + [user_id, mid])
            counts["material_update"] += 1 if m.get("changes") else 0
            detail = "Excel から更新: " + describe_changes(m.get("changes") or {})
        conn.commit()
        if m["action"] == "new" or m.get("changes"):
            db.log_activity(conn, user_id, "excel_import", mid, detail, actor, client)
        if m.get("current_price") is not None or m.get("current_lead") is not None:
            last = db.latest_observation(conn, mid)
            if not last or last["unit_price"] != m.get("current_price") or last["lead_time_days"] != m.get("current_lead"):
                db.add_observation(conn, mid, "manual", m.get("current_price"), m.get("current_lead"), created_by=user_id)
                counts["price"] += 1
    return counts


def describe_changes(changes):
    parts = []
    for k, (old, new) in changes.items():
        label = FIELD_LABELS.get(k, {"owner_id": "担当者", "preferred_supplier_id": "主仕入先", "active": "状態",
                                     "contact_name": "ご担当者名", "email": "メール"}.get(k, k))
        parts.append(f"{label}: {old if old not in (None, '') else '(空)'} → {new if new not in (None, '') else '(空)'}")
    return " / ".join(parts)


def dumps(preview):
    return json.dumps(preview, ensure_ascii=False)
